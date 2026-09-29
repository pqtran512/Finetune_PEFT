"""Export ISBM contradictions review and presentation script to Word."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(__file__).resolve().parent / "ISBM_Contradictions_and_Presentation_Script.docx"


def _add_vi(paragraph, text: str) -> None:
    """Append italic Vietnamese translation on a new line."""
    paragraph.add_run("\n")
    run = paragraph.add_run(text)
    run.italic = True


QA_SECTIONS = [
    (
        "3.1 Method & Contributions",
        [
            (
                "Q1. What is the main contribution of your work compared to existing PEFT studies?",
                "Our main contribution is a Java-focused pipeline, not Python or C++. "
                "We combine QLoRA fine-tuning, unified post-processing, and a compile-aware repair pass. "
                "On HumanEval-Java we reach 80.38% Pass@1, outperforming the Qwen2.5-Coder-32B-Instruct baseline.",
                "Q1. Đóng góp chính của nghiên cứu so với các công trình PEFT hiện có là gì?",
                "Đóng góp chính là pipeline chuyên cho Java, không phải Python hay C++. "
                "Chúng tôi kết hợp fine-tune QLoRA, hậu xử lý thống nhất và bước sửa lỗi dựa trên biên dịch. "
                "Trên HumanEval-Java đạt 80,38% Pass@1, vượt baseline Qwen2.5-Coder-32B-Instruct.",
            ),
            (
                "Q2. Why did you choose CodeLlama-7B and Qwen2.5-Coder-32B?",
                "CodeLlama-7B is a practical starting point: it has a public Java baseline of 29.2% and can be fine-tuned on a 12 GB GPU. "
                "We then moved to Qwen2.5-Coder-32B because it already has strong Java performance at 65.49%, "
                "so we could test whether our pipeline scales to a stronger base model.",
                "Q2. Vì sao chọn CodeLlama-7B và Qwen2.5-Coder-32B?",
                "CodeLlama-7B là điểm khởi đầu thực tế: có baseline Java công khai 29,2% và fine-tune được trên GPU 12 GB. "
                "Sau đó chuyển sang Qwen2.5-Coder-32B vì đã mạnh sẵn ở Java (65,49%), "
                "để kiểm tra pipeline có scale lên mô hình nền mạnh hơn không.",
            ),
            (
                "Q3. Did you use Bayesian Optimization? Slides mention it, but the paper describes manual search.",
                "In the current experiments we did not use Optuna or Bayesian Optimization. "
                "We used a two-phase manual search: Phase 1 on 10% of the data to select hyperparameters, "
                "then Phase 2 on the full dataset. Automatic hyperparameter optimization is part of our future work.",
                "Q3. Có dùng Bayesian Optimization không? Slide có nhắc nhưng paper mô tả tìm kiếm thủ công.",
                "Trong thực nghiệm hiện tại chúng tôi không dùng Optuna hay Bayesian Optimization. "
                "Chúng tôi dùng tìm kiếm thủ công hai giai đoạn: Giai đoạn 1 trên 10% dữ liệu để chọn siêu tham số, "
                "rồi Giai đoạn 2 trên toàn bộ dataset. Tối ưu siêu tham số tự động là hướng tương lai.",
            ),
        ],
    ),
    (
        "3.2 QLoRA & Training",
        [
            (
                "Q4. Why QLoRA instead of full fine-tuning or standard LoRA?",
                "Full fine-tuning needs about 55 to 70 GB of VRAM, which is not feasible on our RTX 3060. "
                "QLoRA stores the base model in 4-bit and trains only small LoRA adapters, "
                "cutting memory use by roughly five to six times while keeping strong results.",
                "Q4. Vì sao dùng QLoRA thay vì full fine-tuning hoặc LoRA thuần?",
                "Full fine-tuning cần khoảng 55–70 GB VRAM, không khả thi trên RTX 3060 của chúng tôi. "
                "QLoRA lưu base model ở 4-bit và chỉ huấn luyện adapter LoRA nhỏ, "
                "giảm bộ nhớ khoảng 5–6 lần mà vẫn giữ kết quả tốt.",
            ),
            (
                "Q5. What hyperparameters did you tune, and how sensitive was the model?",
                "We tuned learning rate, number of epochs, gradient accumulation, and LoRA settings such as rank, alpha, and dropout. "
                "Phase 1 on a 10% subset helped us pick a stable configuration. "
                "Both training and validation loss decreased steadily, with no clear sign of overfitting.",
                "Q5. Đã tinh chỉnh những siêu tham số nào? Mô hình nhạy đến mức nào?",
                "Chúng tôi tinh chỉnh learning rate, số epoch, gradient accumulation và các tham số LoRA như rank, alpha, dropout. "
                "Giai đoạn 1 trên 10% dữ liệu giúp chọn cấu hình ổn định. "
                "Loss train và validation đều giảm đều, không thấy overfitting rõ.",
            ),
            (
                "Q6. Why prefix–target completion instead of instruction/chat format?",
                "HumanEval-Java asks the model to complete a method body from a given signature. "
                "A prefix-target completion format matches that task directly and is closer to real IDE code completion than chat-style prompting.",
                "Q6. Vì sao dùng completion prefix–target thay vì instruction/chat?",
                "HumanEval-Java yêu cầu hoàn thiện thân phương thức từ chữ ký hàm cho sẵn. "
                "Định dạng prefix–target khớp trực tiếp với tác vụ này và gần với code completion trong IDE hơn chat format.",
            ),
        ],
    ),
    (
        "3.3 Evaluation & Results",
        [
            (
                "Q7. Why HumanEval-Java? Does it represent real-world Java development?",
                "We chose HumanEval-Java because it is a standard benchmark with automatic unit tests, "
                "so our results are reproducible and comparable with the BigCode leaderboard. "
                "It does not fully represent enterprise Java, so broader benchmarks are future work.",
                "Q7. Vì sao chọn HumanEval-Java? Có phản ánh phát triển Java thực tế không?",
                "Chúng tôi chọn HumanEval-Java vì là benchmark chuẩn có unit test tự động, "
                "kết quả reproducible và so sánh được với BigCode leaderboard. "
                "Benchmark này chưa phản ánh đủ Java enterprise; mở rộng benchmark là hướng tương lai.",
            ),
            (
                "Q8. Paper reports 65/158 passed, but the dataset has 164 problems. Why?",
                "MultiPL-E defines 164 Java problems, but our evaluation run covers 158 tasks — "
                "the number of problems we successfully generated and evaluated in the final pipeline. "
                "Pass@1 is computed on that evaluated set of 158 problems.",
                "Q8. Paper ghi 65/158 đúng, nhưng dataset có 164 bài. Vì sao?",
                "MultiPL-E định nghĩa 164 bài Java, nhưng lần đánh giá cuối cùng của chúng tôi gồm 158 bài — "
                "số bài đã sinh mã và đánh giá thành công trong pipeline. "
                "Pass@1 được tính trên tập 158 bài đó.",
            ),
            (
                "Q9. How much of 80.38% comes from fine-tuning vs. inference tricks?",
                "From the Qwen baseline of 65.49%: QLoRA adds about 2.9 points to 68.35%, "
                "the prompt directive adds 2.5 points to 70.89%, and the repair pass adds 9.5 points to 80.38%. "
                "So inference-time fixes help a lot, but fine-tuning still improves the base quality.",
                "Q9. 80,38% đến từ fine-tuning hay kỹ thuật ở inference?",
                "Từ baseline Qwen 65,49%: QLoRA cộng ~2,9 điểm lên 68,35%; "
                "prompt directive cộng 2,5 điểm lên 70,89%; repair pass cộng 9,5 điểm lên 80,38%. "
                "Sửa ở inference giúp nhiều, nhưng fine-tune vẫn nâng chất lượng nền.",
            ),
            (
                "Q10. Does repair pass count as Pass@1? Is comparison with single-pass baselines fair?",
                "Yes, because repair is only applied when the first output fails to compile; "
                "already correct solutions are left unchanged. "
                "All models use the same post-processing, and we also report 70.89% without repair as a conservative result.",
                "Q10. Repair pass có được tính vào Pass@1 không? So sánh có công bằng không?",
                "Có, vì repair chỉ chạy khi lần sinh đầu không biên dịch được; "
                "bài đã đúng thì giữ nguyên. "
                "Mọi mô hình dùng chung post-processing, và chúng tôi cũng báo 70,89% không repair để so sánh thận trọng.",
            ),
        ],
    ),
    (
        "3.4 Repair Pass & Post-processing",
        [
            (
                "Q11. How does the repair pass work? Could it introduce new bugs?",
                "After the first generation, we compile the code. If compilation fails, "
                "we read the error message, identify missing helpers or symbols, "
                "and ask the model to generate only those missing parts without changing the existing code. "
                "This limits the risk of breaking solutions that were already correct.",
                "Q11. Repair pass hoạt động thế nào? Có thể gây lỗi mới không?",
                "Sau lần sinh đầu, chúng tôi biên dịch mã. Nếu lỗi biên dịch, "
                "đọc thông báo lỗi, xác định helper/symbol thiếu, "
                "rồi chỉ sinh phần thiếu mà không sửa mã đã có. "
                "Cách này hạn chế làm hỏng các bài đã đúng.",
            ),
            (
                "Q12. Is the unified post-processing module applied fairly to all models?",
                "Yes. The same post-processing module is applied to every model: "
                "it removes Markdown artifacts, balances braces, and extracts helper methods. "
                "That keeps the comparison fair across baselines and fine-tuned models.",
                "Q12. Module hậu xử lý thống nhất có áp dụng công bằng cho mọi mô hình không?",
                "Có. Cùng một module hậu xử lý cho mọi mô hình: "
                "loại Markdown, cân bằng ngoặc nhọn, trích helper. "
                "Nhờ đó so sánh công bằng giữa baseline và mô hình đã fine-tune.",
            ),
        ],
    ),
    (
        "3.5 Limitations & Future Work",
        [
            (
                "Q13. What are the main limitations?",
                "We only measure Pass@1, not code readability or maintainability. "
                "The benchmark is small and not fully representative of real projects. "
                "Repair pass helps compile errors more than logic errors, "
                "and the Qwen experiments require an H100 with 80 GB memory.",
                "Q13. Hạn chế chính là gì?",
                "Chúng tôi chỉ đo Pass@1, chưa đo readability hay maintainability. "
                "Benchmark nhỏ, chưa đại diện đủ dự án thực tế. "
                "Repair pass giúp lỗi biên dịch nhiều hơn lỗi logic; "
                "thí nghiệm Qwen cần H100 80 GB.",
            ),
            (
                "Q14. Can your method generalize to other languages?",
                "Yes, in principle. The same idea — QLoRA fine-tuning, shared post-processing, "
                "and compiler-guided repair — can transfer to other statically typed languages such as C# or Kotlin, "
                "as long as we adapt the dataset, benchmark, and repair rules.",
                "Q14. Phương pháp có generalize sang ngôn ngữ khác không?",
                "Về nguyên tắc là có. Cùng ý tưởng — QLoRA, hậu xử lý chung, sửa lỗi theo compiler — "
                "có thể chuyển sang ngôn ngữ kiểu tĩnh như C# hoặc Kotlin, "
                "nếu thay dataset, benchmark và luật repair cho phù hợp.",
            ),
            (
                "Q15. How does this compare to GitHub Copilot or commercial assistants?",
                "Commercial tools like Copilot are closed-source and hard to reproduce scientifically. "
                "Our work is an open, measurable research pipeline focused on Java and Pass@1 on a public benchmark. "
                "We are not claiming to beat Copilot on all real-world coding tasks.",
                "Q15. So với GitHub Copilot hay công cụ thương mại thì sao?",
                "Copilot và công cụ thương mại là closed-source, khó reproduce khoa học. "
                "Công trình này là pipeline nghiên cứu mở, đo được, tập trung Java và Pass@1 trên benchmark công khai. "
                "Chúng tôi không claim vượt Copilot trên mọi tác vụ lập trình thực tế.",
            ),
        ],
    ),
    (
        "3.6 Tough / Trap Questions",
        [
            (
                "Q16. If repair pass gives +9.5 pp, why bother with fine-tuning?",
                "Because repair only fixes compilation errors, not wrong logic. "
                "Fine-tuning still raises the baseline from 65.49% to 68.35% and improves overall code quality. "
                "The best results come from combining both, not from repair alone.",
                "Q16. Repair pass đã +9,5 điểm rồi, còn fine-tune làm gì?",
                "Vì repair chỉ sửa lỗi biên dịch, không sửa logic sai. "
                "Fine-tune vẫn nâng baseline từ 65,49% lên 68,35% và cải thiện chất lượng tổng thể. "
                "Kết quả tốt nhất đến từ kết hợp cả hai, không phải chỉ repair.",
            ),
            (
                "Q17. Your 7B beats 34B — is that because of data leakage or benchmark contamination?",
                "We do not train directly on HumanEval-Java. "
                "Training data comes from Evol-Instruct-Code with filtering and decontamination against benchmark overlap. "
                "Evaluation is on held-out HumanEval-Java tasks, so the gain comes from domain adaptation, not leakage.",
                "Q17. 7B vượt 34B — có phải do rò rỉ dữ liệu hoặc benchmark contamination?",
                "Chúng tôi không train trực tiếp trên HumanEval-Java. "
                "Dữ liệu train từ Evol-Instruct-Code, có lọc và decontaminate trùng benchmark. "
                "Đánh giá trên tập held-out HumanEval-Java, nên cải thiện đến từ thích ứng miền, không phải leakage.",
            ),
            (
                "Q18. What is the computational cost and environmental impact?",
                "With QLoRA, CodeLlama-7B was fine-tuned on one RTX 3060 with 12 GB VRAM. "
                "Full fine-tuning of the same 7B model typically needs about 55 to 70 GB of VRAM — "
                "roughly five to six times more — so it would not fit on our GPU at all. "
                "For Qwen2.5-Coder-32B, QLoRA 4-bit training took about four days on one H100 80 GB GPU. "
                "Full fine-tuning of a 32B model would need far more memory — often over 120 GB or multiple GPUs — "
                "which was not feasible for us. In short, QLoRA did not just reduce cost; it made these experiments possible on the hardware we had.",
                "Q18. Chi phí tính toán và tác động môi trường thế nào?",
                "Với QLoRA, CodeLlama-7B fine-tune trên một RTX 3060 12 GB VRAM. "
                "Full fine-tuning cùng mô hình 7B thường cần khoảng 55–70 GB VRAM — "
                "nhiều hơn khoảng 5–6 lần — nên không chạy được trên GPU của chúng tôi. "
                "Với Qwen2.5-Coder-32B, QLoRA 4-bit mất khoảng 4 ngày trên một H100 80 GB. "
                "Full fine-tuning mô hình 32B cần bộ nhớ lớn hơn nhiều — thường trên 120 GB hoặc nhiều GPU — "
                "không khả thi với chúng tôi. Tóm lại, QLoRA không chỉ rẻ hơn mà còn giúp thí nghiệm có thể thực hiện được trên phần cứng hiện có.",
            ),
        ],
    ),
]


FALLBACK_TIPS = [
    (
        "You don't know the answer",
        "That is a good question. We have not tested that yet, but it is a direction we plan to explore in future work.",
        "Không biết câu trả lời",
        "Đây là câu hỏi hay. Chúng tôi chưa thử nghiệm điểm đó, nhưng đó là hướng dự định làm trong tương lai.",
    ),
    (
        "Asked about Bayesian Optimization",
        "We used two-phase manual hyperparameter selection in the current study. Automatic Bayesian Optimization is planned for future work.",
        "Bị hỏi về Bayesian Optimization",
        "Trong nghiên cứu hiện tại chúng tôi dùng chọn siêu tham số thủ công hai giai đoạn. Bayesian Optimization tự động là hướng tương lai.",
    ),
    (
        "Asked about 161 vs 164 vs 158",
        "MultiPL-E has 164 Java problems. Our reported Pass@1 is on 158 problems that were generated and evaluated in the final run.",
        "Bị hỏi về 161 / 164 / 158",
        "MultiPL-E có 164 bài Java. Pass@1 báo cáo được tính trên 158 bài đã sinh mã và đánh giá trong lần chạy cuối.",
    ),
    (
        "Asked whether repair pass is fair",
        "Repair is applied only after a compile error, and we also report 70.89% without repair for a conservative comparison.",
        "Bị hỏi repair pass có công bằng không",
        "Repair chỉ chạy sau lỗi biên dịch, và chúng tôi cũng báo 70,89% không repair để so sánh thận trọng.",
    ),
]


def add_qa_section(doc: Document) -> None:
    doc.add_heading("Part 3 — Potential Committee Questions & Answers", level=1)
    doc.add_paragraph(
        "Short answers ready to speak (about 30–45 seconds each). "
        "Questions are likely in English at ISBM. "
        "Vietnamese translations below each item are in italics for rehearsal."
    )
    doc.add_paragraph()

    for section_title, items in QA_SECTIONS:
        doc.add_heading(section_title, level=2)
        for question_en, answer_en, question_vi, answer_vi in items:
            p_q = doc.add_paragraph()
            run_q = p_q.add_run(question_en)
            run_q.bold = True
            _add_vi(p_q, question_vi)

            p_a = doc.add_paragraph()
            run_a = p_a.add_run("Answer: ")
            run_a.bold = True
            p_a.add_run(answer_en)
            _add_vi(p_a, f"Trả lời: {answer_vi}")

            doc.add_paragraph()

    doc.add_heading("3.7 Quick Fallback Answers", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Situation"
    table.rows[0].cells[1].text = "Answer"
    for sit_en, resp_en, sit_vi, resp_vi in FALLBACK_TIPS:
        row = table.add_row().cells
        row[0].text = sit_en
        cell_b = row[1]
        cell_b.text = resp_en
        # Re-add with italic VI — table cells need paragraph runs
        cell_b.text = ""
        p = cell_b.paragraphs[0]
        p.add_run(resp_en)
        _add_vi(p, f"Trả lời: {resp_vi}")
        cell_a = row[0]
        cell_a.text = ""
        pa = cell_a.paragraphs[0]
        pa.add_run(sit_en)
        _add_vi(pa, sit_vi)


def main():
    doc = Document()

    title = doc.add_heading("ISBM Presentation — Content Review & Speaker Script", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("Paper: Tran_Springer_Template.docx")
    doc.add_paragraph("Slides: ISBM_Java_Code_Generation_Slides.pptx")
    doc.add_paragraph("Prepared for: Pham Quynh Tran — ISBM Conference (~10 min)")
    doc.add_paragraph()

    # --- Part 1 ---
    doc.add_heading("Part 1 — Contradictions Between Paper and Slides", level=1)
    doc.add_heading("1.1 High Priority (fix before presenting)", level=2)

    table1 = doc.add_table(rows=1, cols=4)
    table1.style = "Table Grid"
    hdr = table1.rows[0].cells
    hdr[0].text = "#"
    hdr[1].text = "Slides say"
    hdr[2].text = "Paper says"
    hdr[3].text = "Note"

    rows_high = [
        (
            "1",
            "Bayesian Optimization (Optuna TPE) is a core contribution "
            "(title, slide 4, motivation, conclusion)",
            "BO only mentioned in Related Work as a general method. Experiments use "
            "two-phase manual hyperparameter search on 10% data. Future work states "
            "automatic hyperparameter optimization could be applied.",
            "Major mismatch. Slides imply BO was done; paper lists it as future work.",
        ),
        (
            "2",
            "Title: Improving Java Code Generation with Parameter-Efficient Fine-Tuning",
            "Title: Java Code Generation using Parameter Tuning",
            "Branding inconsistency — align before conference.",
        ),
        (
            "3",
            "HumanEval-Java: 161 problems",
            "Dataset described as 164 problems; results report 65/158 Pass@1.",
            "Use 164 (benchmark size) or 158 (evaluated), not 161.",
        ),
        (
            "4",
            "Compares to CodeLlama-34B-Instruct (41.53%)",
            "Table 2 only has CodeLlama-34B (40.19%); text compares to 34B, not Instruct.",
            "Slide adds a model not in the paper table.",
        ),
    ]
    for row_data in rows_high:
        row = table1.add_row().cells
        for i, val in enumerate(row_data):
            row[i].text = val

    doc.add_paragraph()
    doc.add_heading("1.2 Medium Priority (slides more detailed than paper)", level=2)

    table2 = doc.add_table(rows=1, cols=2)
    table2.style = "Table Grid"
    table2.rows[0].cells[0].text = "On slides"
    table2.rows[0].cells[1].text = "In paper"
    rows_med = [
        ("NEFTune, cosine LR, decontaminate HumanEval", "Not mentioned"),
        ("Class Problem wrapper ~80%", "Not mentioned"),
        ("Magicoder (in old pipeline)", "Only Evol-Instruct-Code-80k-v1"),
        ("run_pipeline.py, 7 automated steps", "Not mentioned"),
        ("Framework name PEFT-JAC (Section 3)", "Slides do not use this name"),
    ]
    for a, b in rows_med:
        row = table2.add_row().cells
        row[0].text = a
        row[1].text = b

    doc.add_paragraph()
    doc.add_heading("1.3 No contradiction — scope difference only", level=2)

    table3 = doc.add_table(rows=1, cols=2)
    table3.style = "Table Grid"
    table3.rows[0].cells[0].text = "Content"
    table3.rows[0].cells[1].text = "Note"
    rows_ok = [
        (
            "Flask Prototype (slide 8)",
            "Slide-only demo; not in paper — OK if presented as supplementary.",
        ),
        (
            "Unified post-processing module (paper contribution ii)",
            "Paper emphasizes this; slides only mention post-processing briefly.",
        ),
        (
            "Pass@1 numbers",
            "Match: 29.20% → 41.14%; Qwen 65.49% → 68.35% → 70.89% → 80.38%; repair +9.5 pp.",
        ),
        ("Hardware", "RTX 3060 12GB / H100 80GB — match."),
        ("Training params (Table 1)", "LR 3e-4/5e-5, epochs 3/2, grad accum 32/16 — match."),
    ]
    for a, b in rows_ok:
        row = table3.add_row().cells
        row[0].text = a
        row[1].text = b

    doc.add_paragraph()
    doc.add_heading("1.4 Recommended slide fixes", level=2)
    fixes = [
        'Remove or downgrade Bayesian Optimization — replace with "two-phase hyperparameter selection".',
        'Fix 161 → 164 (or note "158 evaluated problems").',
        "Remove CodeLlama-34B-Instruct from results table, or add it to the paper first.",
        "Align slide title with Springer paper title if slides accompany the paper.",
    ]
    for item in fixes:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_page_break()

    # --- Part 2 ---
    doc.add_heading("Part 2 — Presentation Script (English, ~10 minutes)", level=1)
    doc.add_paragraph(
        "Total speaking time: approximately 9:30–10:00 + Q&A. "
        "Speak slowly and clearly on results slides."
    )
    doc.add_paragraph()

    slides = [
        (
            "Slide 1 — Title (~40 seconds)",
            """Good morning/afternoon everyone. Thank you for being here.

My name is Pham Quynh Tran from Ho Chi Minh City University of Technology.

Today I will present our work on improving Java code generation using parameter-efficient fine-tuning.

We combine QLoRA, careful hyperparameter tuning, and inference-time repair to boost code quality on the HumanEval-Java benchmark — using both CodeLlama-7B and Qwen2.5-Coder-32B.""",
        ),
        (
            "Slide 2 — Outline (~25 seconds)",
            """Here is the outline.

I will start with the motivation, briefly review related work, then describe our pipeline and method.

After that I will explain evaluation, present results on both models, show a short demo prototype, and conclude with future directions.""",
        ),
        (
            "Slide 3 — Motivation (~1:00)",
            """Java remains widely used in industry, but general-purpose code LLMs often underperform on Java compared to languages like Python.

One reason is training data imbalance and Java's strict syntax and compilation requirements.

Full fine-tuning of large models is also expensive — often impractical on a single consumer GPU.

At the same time, hyperparameters such as learning rate, LoRA rank, and batch size strongly affect generation quality.

Our CodeLlama-7B baseline scores only 29.2% Pass@1 on HumanEval-Java.

So our goals are: specialize models for Java with PEFT/QLoRA, tune hyperparameters systematically, build a reproducible pipeline, and improve compile-time correctness through prompt and repair strategies.""",
        ),
        (
            "Slide 4 — Related Work (~50 seconds)",
            """Prior PEFT studies for code generation focus mainly on Python or C++, not Java.

Li et al. showed that LoRA and IA³ can improve secure code generation for C/C++ more efficiently than full fine-tuning.

Weyssow et al. surveyed PEFT for code LLMs on Python and found LoRA and QLoRA most effective.

Our work fills the gap: a Java-specific pipeline with QLoRA, structured hyperparameter selection, and compile-aware repair at inference time.""",
        ),
        (
            "Slide 5 — Proposed Pipeline (~1:00)",
            """Our system follows a modular pipeline.

Input prompts go through preprocessing — template formatting and tokenization.

The core LLM with QLoRA adapters generates Java code.

Post-processing cleans markdown artifacts, balances braces, and extracts helper methods — this unified module ensures fair comparison across models.

Finally, we evaluate with javac and unit tests.

We started with CodeLlama-7B and scaled to Qwen2.5-Coder-32B.

After training, LoRA weights are merged into the base model for simpler deployment.

The full workflow covers dataset building, hyperparameter experiments, QLoRA training, merging, inference, and Pass@1 evaluation.""",
        ),
        (
            "Slide 6 — Method: QLoRA & Hyperparameter Tuning (~1:15)",
            """Three key components.

First, QLoRA: the base model is frozen in 4-bit NF4; we train low-rank adapters with W equals W₀ plus BA, targeting attention and MLP layers. This fits on 12 GB or 80 GB GPUs.

Second, hyperparameter tuning: we use a two-phase strategy. Phase 1 runs experiments on a 10% subset to find stable settings. Phase 2 trains on the full dataset with adjusted batch size and gradient accumulation.

Third, training format: data comes from Evol-Instruct-Code, filtered to about 13,600 Java samples, mapped to prefix–target completion — prefix masked, target trained — split 90/10 for train and validation.""",
        ),
        (
            "Slide 7 — Evaluation (~50 seconds)",
            """We evaluate on HumanEval-Java from MultiPL-E — 164 programming problems with unit tests.

We report Pass@1 with greedy decoding, aligned with the BigCode leaderboard.

The pipeline: generate code → post-process → compile with javac → run tests.

Outcomes are classified as PASSED, COMPILE_ERROR, or FAILED.

Error analysis guided our improvements: compile errors are addressed with prompt directives and repair pass; logic failures need a stronger base model or richer training data.""",
        ),
        (
            "Slide 8 — Results: CodeLlama-7B (~1:10)",
            """On a single RTX 3060 with 12 GB VRAM, our fine-tuned CodeLlama-7B reaches 41.14% Pass@1.

That is +11.9 percentage points over the 29.2% baseline.

It outperforms CodeLlama-13B and 34B, and approaches much larger models — showing that PEFT on limited hardware can beat bigger frozen models.

Training and validation loss both decreased steadily with no clear overfitting — the model generalized well.

This exceeded our initial target of beating the published 29.2% baseline.""",
        ),
        (
            "Slide 9 — Results: Qwen2.5-Coder-32B (~1:30)",
            """We then scaled to Qwen2.5-Coder-32B on an H100 80 GB, reusing the same pipeline.

Starting from 65.49%, QLoRA SFT raises Pass@1 to 68.35%.

Adding a prompt directive — Java comments enforcing complete methods — brings it to 70.89%, without retraining.

The repair pass gives the largest gain: 80.38%, up 9.5 points. It reads compiler errors, generates missing helper methods only, and keeps already-correct code unchanged.

The final score beats Qwen2.5-Coder-32B-Instruct at 73.69% by 6.7 points — showing that domain-specific fine-tuning plus targeted inference repair can outperform general instruction tuning.""",
        ),
        (
            "Slide 10 — Demo Prototype (~45 seconds)",
            """To bridge research and practice, we built a Flask web demo with CodeLlama-7B plus LoRA.

Users can input natural language, a function signature, or a full Java class.

A prompt enhancer applies Java templates and Javadoc rules.

Generated code is verified with javac and can be run in a sandbox.

This is a supplementary demo — the core contributions are in the paper's training and evaluation pipeline.""",
        ),
        (
            "Slide 11 — Conclusion & Future Work (~50 seconds)",
            """To summarize:

We built an end-to-end Java PEFT pipeline with hyperparameter tuning and compile-aware repair.

Results: 41.14% on CodeLlama-7B with consumer GPU, up to 80.38% on Qwen2.5-Coder-32B.

Key insight: analyzing failure modes and fixing structural errors at inference can be more cost-effective than blind retraining.

Future work includes automatic hyperparameter optimization, larger Java-specific datasets, multi-benchmark evaluation, and code quality beyond Pass@1 — readability and maintainability.""",
        ),
        (
            "Slide 12 — Thank You (~15 seconds)",
            """Thank you for your attention.

I am happy to take your questions.""",
        ),
    ]

    for heading, body in slides:
        doc.add_heading(heading, level=2)
        for para in body.strip().split("\n\n"):
            doc.add_paragraph(para)
        doc.add_paragraph()

    doc.add_page_break()
    add_qa_section(doc)

    doc.add_heading("Pre-conference checklist", level=2)
    checklist = [
        "Decide: align slides with paper (remove BO) or update paper if Optuna BO was actually used.",
        "Fix 161 → 164/158 on Evaluation slide.",
        "Align title with Springer paper.",
        'If asked about BO: answer "two-phase manual search; automatic BO is future work" — matches paper.',
    ]
    for item in checklist:
        doc.add_paragraph(item, style="List Bullet")

    doc.save(str(OUT))
    print(f"Saved: {OUT}")


if __name__ == "__main__":
    main()
